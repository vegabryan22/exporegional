import json
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from pypdf import PdfReader
from app.extensions import db
from app.models.assignment import Assignment
from app.models.project import Project
from app.models.institution import Institution
from app.models.system_setting import SystemSetting
from app.services.expo_progress_service import VENUE_SETTING
from tests import test_expo_attendance as attendance_fixture


class UsherReportTest(unittest.TestCase):
    def setUp(self):
        self.fixture = attendance_fixture.ExpoAttendanceTest()
        self.fixture.setUp()
        self.app = self.fixture.app
        self.app.secret_key = 'test'
        first = self.fixture.project
        first.institution.name = 'CTP Zeta'
        school = Institution(code='A', name='CTP Águila', responsible_name='Persona', responsible_email='a@school.test')
        db.session.add(school); db.session.flush()
        second = Project(title='Proyecto inglés', team_name='Equipo dos', representative_name='Persona',
                         representative_email='p@test', description='Descripción', category='steam',
                         institution=school, is_active=True)
        db.session.add(second); db.session.flush()
        judge = self.fixture.judges[0]
        judge.full_name = 'Juez de inglés'
        judge.can_evaluate_documentation = False
        judge.can_evaluate_exposition = False
        judge.can_evaluate_english = True
        self.assignments = [
            Assignment(judge=self.fixture.judges[1], project=first, status=Assignment.STATUS_CONFIRMED,
                       can_evaluate_exposition=True, can_evaluate_documentation=False),
            Assignment(judge=judge, project=second, status=Assignment.STATUS_CONFIRMED,
                       can_evaluate_exposition=False, can_evaluate_documentation=False, can_evaluate_english=True),
            Assignment(judge=self.fixture.judges[2], project=second, status=Assignment.STATUS_DRAFT,
                       can_evaluate_exposition=True),
            Assignment(judge=self.fixture.judges[3], project=first, status=Assignment.STATUS_CONFIRMED,
                       can_evaluate_documentation=True, can_evaluate_exposition=False),
        ]
        db.session.add_all(self.assignments)
        config = {'venues':[{'id':'a4','name':'P1-A4','responsible':'Carlos Ticas'}],
                  'projects':{str(second.id):'a4'}}
        SystemSetting.set_value(VENUE_SETTING,json.dumps(config)); db.session.flush()
        self.context = {'projects':[first,second], 'assignments':self.assignments, 'category_map':{'steam':'STEAM'}}

    def tearDown(self):
        self.fixture.tearDown()

    def exports(self):
        from app.controllers import admin_controller as controller
        with self.app.test_request_context('/'), patch.object(controller,'_base_context',return_value=self.context), patch.object(controller,'_institution_name',return_value='ExpoTécnica Regional'):
            xlsx = controller.exposition_usher_report_excel.__wrapped__()
            pdf = controller.exposition_usher_report_pdf.__wrapped__()
            xlsx.direct_passthrough = False; pdf.direct_passthrough = False
            return xlsx.get_data(), pdf.get_data()

    def test_rows_sort_by_project_school_and_include_english_only(self):
        from app.controllers.admin_controller import _build_exposition_usher_report_rows
        rows = _build_exposition_usher_report_rows(self.context)
        self.assertEqual(['CTP Águila','CTP Zeta'],[row['school'] for row in rows])
        self.assertEqual('Inglés',rows[0]['evaluation'])
        self.assertEqual('P1-A4',rows[0]['location'])
        self.assertEqual('Carlos Ticas',rows[0]['responsible'])
        self.assertEqual('Colegio 2',rows[0]['judge_school'])
        self.assertEqual('Pendiente de asignar',rows[1]['location'])

    def test_exports_have_real_venues_school_columns_and_followup_validation(self):
        xlsx, pdf = self.exports()
        sheet = load_workbook(BytesIO(xlsx)).active
        self.assertEqual('Colegio del proyecto',sheet['A5'].value)
        self.assertEqual('Colegio que inscribió al juez',sheet['G5'].value)
        self.assertEqual('CTP Águila',sheet['A6'].value)
        self.assertEqual('P1-A4',sheet['D6'].value)
        self.assertEqual('Inglés',sheet['F6'].value)
        self.assertEqual('L6:L7',str(sheet.data_validations.dataValidation[0].sqref))
        self.assertEqual('A5:L7',sheet.tables['GuiaEdecanes'].ref)
        self.assertEqual('C6',sheet.freeze_panes)
        text = '\n'.join(page.extract_text() for page in PdfReader(BytesIO(pdf)).pages)
        self.assertIn('P1-A4',text)
        self.assertIn('Juez de inglés',text)
        self.assertIn('Colegio 2',text)
        self.assertLess(text.index('CTP Águila'),text.index('CTP Zeta'))

    def test_pdf_repeats_headers_on_multiple_pages(self):
        self.context['assignments'] = self.assignments * 35
        _, pdf = self.exports()
        pages = PdfReader(BytesIO(pdf)).pages
        self.assertGreater(len(pages),1)
        for page in pages:
            text = page.extract_text()
            self.assertIn('Colegio del proyecto',text)
            self.assertIn('Recinto',text)
            self.assertIn('Página',text)


if __name__ == '__main__':
    # Synthetic QA output only; not an export of production judge data.
    test = UsherReportTest(); test.setUp()
    try:
        xlsx, pdf = test.exports()
        out = Path('tmp/usher-report-qa'); out.mkdir(parents=True, exist_ok=True)
        (out/'guide.xlsx').write_bytes(xlsx); (out/'guide.pdf').write_bytes(pdf)
        print('Synthetic export QA files:',out)
    finally:
        test.tearDown()
